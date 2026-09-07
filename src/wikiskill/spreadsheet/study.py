"""Opt-in, isolated, single-round Spreadsheet development study.

This controller owns immutable inputs, scoring and gates. The isolated runtime
owns model execution and OS boundaries. No legacy runner, TEST or next round is
reachable here. Stable UID prefixes are a software smoke selection, not a
random scientific sample. Only completed native calls can become score rows.
"""
from __future__ import annotations

from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from contextlib import contextmanager
from datetime import datetime, timezone
import fcntl
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import zipfile
from xml.etree.ElementTree import ParseError

from wikiskill.benchmarks import spreadsheet
from wikiskill.jsonl import read_jsonl
from wikiskill.paper_alignment import contracts, evidence

SCHEMA = "wikiskill.spreadsheet.study.v1"
PACKAGE = Path(__file__).resolve().parents[1]
PROMPTS = PACKAGE / "resources/paper_alignment/prompts"
SEED = 2026090705


class StudyError(RuntimeError):
    """Study stopped; inspect its preserved artifacts before resuming."""


class IntegrityError(StudyError):
    """A frozen input, execution or derivation no longer matches."""


def _now():
    return datetime.now(timezone.utc).isoformat()


def _sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _text_sha(value):
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n"


def _save_text(path, text):
    path = Path(path)
    if path.exists():
        if path.is_symlink() or path.read_text(encoding="utf-8") != text:
            raise IntegrityError(f"Frozen artifact differs: {path}")
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        handle.write(text)


def _save(path, value):
    _save_text(path, _json(value))


def _tree(root):
    root = Path(root)
    result = {}
    for path in sorted(root.rglob("*")):
        if path.is_symlink():
            raise IntegrityError(f"Symlink in sealed tree: {path}")
        if path.is_file():
            result[str(path.relative_to(root))] = _sha(path)
    return result


def _check_tree(root, expected):
    if _tree(root) != expected:
        raise IntegrityError(f"Sealed tree membership or contents changed: {root}")


def _local(root, relative):
    path = Path(root) / relative
    if not path.resolve().is_relative_to(Path(root).resolve()):
        raise IntegrityError("Artifact path escapes study root")
    return path


@contextmanager
def _lock(root):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    with (root / "study.lock").open("a+") as handle:
        try:
            fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as exc:
            raise StudyError("Another study controller holds the writer lock") from exc
        try:
            yield
        finally:
            fcntl.flock(handle, fcntl.LOCK_UN)


def _state(root, phase, **fields):
    """Append an event; state.json is only the replaceable latest projection."""
    events = root / "events"
    events.mkdir(exist_ok=True)
    index = len(list(events.glob("*.json"))) + 1
    value = {"schema_version": SCHEMA, "phase": phase, "updated_at": _now(),
             "active": phase not in ("prepared", "complete", "needs_attention"), **fields}
    _save(events / f"{index:06}.json", value)
    pending = root / "state.json.tmp"
    pending.write_text(_json(value), encoding="utf-8")
    os.replace(pending, root / "state.json")
    return value


def _plain(value):
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if hasattr(value, "text"):
        return {"formula": value.text, "ref": getattr(value, "ref", None)}
    return str(value)


def _preview(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=False)
    try:
        return {sheet.title: [[{"cell": cell.coordinate, "value": _plain(cell.value)} for cell in row]
                             for row in sheet.iter_rows(min_row=1, max_row=min(6, sheet.max_row),
                                                        max_col=min(16, sheet.max_column))]
                for sheet in workbook.worksheets}
    finally:
        workbook.close()


def _source_bindings():
    names = [Path(__file__), Path(spreadsheet.__file__), Path(contracts.__file__), Path(evidence.__file__),
             PACKAGE / "jsonl.py", PACKAGE / "codex_identity.py", *sorted((PACKAGE / "isolated").glob("*.py")),
             *sorted((PACKAGE / "resources/isolated").glob("*.json"))]
    return {str(path.relative_to(PACKAGE)): _sha(path) for path in names}


def _split_ids(path):
    items = _read(path)
    if not isinstance(items, list):
        raise IntegrityError("Split must be a JSON list")
    ids = [str(item["uid"]) if isinstance(item, dict) else str(item) for item in items]
    if len(set(ids)) != len(ids) or any(not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*", uid)
                                       or uid in (".", "..") for uid in ids):
        raise IntegrityError("Split contains duplicate or unsafe UIDs")
    return sorted(ids)


def _codex_binding(*, with_version=True):
    executable = shutil.which("codex")
    if not executable:
        raise IntegrityError("Codex executable is unavailable on PATH")
    executable = Path(executable).resolve()
    result = {"path": str(executable), "sha256": _sha(executable)}
    if with_version:
        process = subprocess.run([str(executable), "--version"], text=True, capture_output=True,
                                 check=True, timeout=30)
        result["version"] = process.stdout.strip()
    return result


def prepare(root: Path, *, data: Path, split_dir: Path, libreoffice_app: Path,
            model: str = "gpt-5.6-luna", effort: str = "high", train_limit: int = 8,
            val_limit: int = 4, workers: int = 2, timeout: int = 1800) -> dict:
    """Freeze a fresh study; selected data are copied, no model is called.

    Limits are intentionally capped at 8 TRAIN and 4 VAL (16 inference calls,
    two learning calls maximum). At least four TRAIN cases support the paper
    Proposer contract. ``root`` must be new/empty; use ``run`` to resume.
    """
    root, data, split_dir = Path(root).resolve(), Path(data).resolve(), Path(split_dir).resolve()
    libreoffice_app = Path(libreoffice_app).resolve()
    if not (4 <= train_limit <= 8 and 1 <= val_limit <= 4 and 1 <= workers <= 2 and timeout > 0):
        raise ValueError("Require 4 <= train_limit <= 8, 1 <= val_limit <= 4, workers 1..2, timeout > 0")
    if (model, effort) != ("gpt-5.6-luna", "high"):
        raise ValueError("This isolated study supports only gpt-5.6-luna/high")
    soffice = libreoffice_app / "Contents/MacOS/soffice"
    if not soffice.is_file():
        raise ValueError("libreoffice_app must contain Contents/MacOS/soffice")
    if root.exists() and any(root.iterdir()):
        raise IntegrityError("Prepare requires an empty study directory; existing studies are immutable")
    codex = _codex_binding()
    ids = {split: _split_ids(split_dir / f"{split}.json") for split in ("train", "val")}
    if set(ids["train"]) & set(ids["val"]):
        raise IntegrityError("TRAIN and VAL overlap")
    selected = {split: ids[split][:limit] for split, limit in (("train", train_limit), ("val", val_limit))}
    if len(selected["train"]) != train_limit or len(selected["val"]) != val_limit:
        raise IntegrityError("Split has fewer cases than the explicitly requested smoke limit")
    cases = spreadsheet.load_cases(data)
    if len({case.uid for case in cases}) != len(cases):
        raise IntegrityError("Dataset contains duplicate UIDs")
    pool = {case.uid: case for case in cases}
    if any(uid not in pool for values in selected.values() for uid in values):
        raise IntegrityError("Selected split UID missing or quarantined in dataset")
    with _lock(root):
        frozen = root / "frozen"
        frozen.mkdir()
        sources = {str(data / "dataset.json"): _sha(data / "dataset.json")}
        rows = {"train": [], "val": []}
        for split in ("train", "val"):
            source_split = split_dir / f"{split}.json"
            sources[str(source_split)] = _sha(source_split)
            shutil.copy2(source_split, frozen / f"{split}-source.json")
            for uid in selected[split]:
                case = pool[uid]
                task_dir = Path(case.task_dir).resolve()
                if not task_dir.is_relative_to(data):
                    raise IntegrityError("Dataset task directory escapes data root")
                inputs = sorted(task_dir.glob("*init*.xlsx"))
                golds = sorted(task_dir.glob("*golden*.xlsx"))
                if len(inputs) != 1 or len(golds) != 1:
                    raise IntegrityError(f"Ambiguous/missing workbooks for {uid}")
                directory = frozen / "data" / split / uid
                directory.mkdir(parents=True)
                for origin, name in ((inputs[0], "input.xlsx"), (golds[0], "reference.xlsx")):
                    if origin.is_symlink() or not origin.resolve().is_relative_to(data):
                        raise IntegrityError("Workbook is outside the dataset")
                    sources[str(origin)] = _sha(origin)
                    shutil.copy2(origin, directory / name)
                reference = directory / "reference.xlsx"
                score, matched, total = spreadsheet.score_workbook(reference, reference, case.answer_sheet, case.answer_position)
                if score != 1 or total <= 0 or matched != total:
                    raise IntegrityError("Reference region cannot be scored")
                rows[split].append({"uid": uid, "split": split, "question": case.instruction.replace("\r\n", "\n").replace("\r", "\n"),
                                    "instruction_type": case.instruction_type, "answer_position": case.answer_position,
                                    "answer_sheet": case.answer_sheet, "preview": _preview(directory / "input.xlsx"),
                                    "input_path": str((directory / "input.xlsx").relative_to(root)),
                                    "reference_path": str(reference.relative_to(root)), "total_cells": total})
        _save(frozen / "cases.json", rows)
        _save(frozen / "data-source-hashes.json", sources)
        for mode in ("spreadsheet", "maintainer", "proposer"):
            destination = frozen / "prompts" / f"{mode}.paper.md"
            destination.parent.mkdir(exist_ok=True)
            shutil.copy2(PROMPTS / destination.name, destination)
        _save(frozen / "initial-wiki.json", _empty_wiki())
        _save(frozen / "initial-skills.json", {})
        protocol = {"schema_version": SCHEMA, "created_at": _now(), "root": str(root), "domain": "spreadsheet",
                    "model": model, "effort": effort, "workers": workers, "timeout_seconds": timeout,
                    "libreoffice_app": str(libreoffice_app), "soffice_sha256": _sha(soffice),
                    "codex": codex,
                    "environment": {"python": sys.version.split()[0], "openpyxl": importlib.metadata.version("openpyxl")},
                    "selection": {"method": "stable UID sort, first N within existing TRAIN/VAL; software smoke only",
                                  "train": selected["train"], "val": selected["val"]},
                    "sampling": {"seed": SEED, "failure_max": 5, "success_max": 3,
                                 "shortage_policy": "use all available in each class; never relabel, replace or resample",
                                 "trace_character_cap": 15000},
                    "budget": {"inference_max": train_limit + 2 * val_limit, "optimizer_max": 2,
                               "train": train_limit, "s0_val": val_limit, "candidate_val_max": val_limit, "test": 0},
                    "sequence": ["fresh_s0_val", "fresh_s0_train", "maintainer", "proposer", "candidate_val_or_no_action", "freeze"],
                    "gate": "ACCEPT iff complete candidate VAL correct count > fresh S0 VAL correct count; ties REJECT",
                    "failure_policy": "Stop dispatch, drain already running calls, preserve raw attempts; no automatic model retries",
                    "scorer": "Existing score_workbook unchanged; data_only cached-value, exact all target cells; missing output = 0",
                    "automatic_next_stage": "none; no K4, TEST or follow-up run",
                    "scope": "Small development end-to-end verification; not held-out generalization or statistical evidence",
                    "source_bindings": _source_bindings(),
                    "resource_bindings": {str(p.relative_to(PACKAGE)): _sha(p) for p in [PROMPTS / f"{mode}.paper.md" for mode in ("spreadsheet", "maintainer", "proposer")]},
                    "frozen_files": _tree(frozen)}
        _save(root / "protocol.json", protocol)
        _save(root / "protocol-lock.json", {"sha256": _sha(root / "protocol.json")})
        _state(root, "prepared", counts={split: len(values) for split, values in rows.items()}, budget=protocol["budget"])
    return status(root)


def _empty_wiki():
    return {"index.md": "# Wiki index\n\nNo patterns have been learned.\n", "log.md": "# Evolution log\n",
            "skill-impact.md": "# Skill impact\n\nNo candidates evaluated. Active skills are empty.\n"}


def _verify_protocol(root):
    protocol = _read(root / "protocol.json")
    if _read(root / "protocol-lock.json") != {"sha256": _sha(root / "protocol.json")}:
        raise IntegrityError("Frozen protocol changed")
    if protocol.get("schema_version") != SCHEMA or protocol.get("root") != str(root.resolve()):
        raise IntegrityError("Protocol identity/root mismatch")
    _check_tree(root / "frozen", protocol["frozen_files"])
    for relative, digest in {**protocol["source_bindings"], **protocol["resource_bindings"]}.items():
        if _sha(_local(PACKAGE, relative)) != digest:
            raise IntegrityError(f"Bound source/resource changed: {relative}")
    if _sha(Path(protocol["libreoffice_app"]) / "Contents/MacOS/soffice") != protocol["soffice_sha256"]:
        raise IntegrityError("Bound LibreOffice binary changed")
    if _codex_binding(with_version=False) != {key: protocol["codex"][key] for key in ("path", "sha256")}:
        raise IntegrityError("Bound Codex executable changed")
    current_environment = {"python": sys.version.split()[0], "openpyxl": importlib.metadata.version("openpyxl")}
    if current_environment != protocol["environment"]:
        raise IntegrityError("Scoring/runtime environment version changed")
    return protocol


def _native(archive):
    if not (archive / "native-complete.json").is_file():
        raise IntegrityError(f"Incomplete native call cannot be scored: {archive}")
    native = _read(archive / "native-complete.json")
    for relative, digest in native["artifacts"].items():
        if _sha(_local(archive, relative)) != digest:
            raise IntegrityError("Native completion artifact changed")
    return native


def _audit(archive, mode, protocol):
    from wikiskill.isolated.runtime import audit, verify_native_complete
    verify_native_complete(archive)
    audit(archive, mode, model=protocol["model"], effort=protocol["effort"])


def _record(path, row):
    _save(path, row)
    _save(path.with_suffix(".seal.json"), {"sha256": _sha(path)})


def _check_record(path):
    if _read(path.with_suffix(".seal.json")) != {"sha256": _sha(path)}:
        raise IntegrityError(f"Sealed record changed: {path}")
    row = _read(path)
    _check_tree(path.parent / "runtime", row["artifacts"])
    _native(path.parent / "runtime")
    return row


def _score(root, case, archive):
    """Task output failures score 0; infrastructure failures raise, never 0."""
    output = archive / "output.xlsx"
    if not output.exists():
        return {"score": 0.0, "matched_cells": 0, "total_cells": case["total_cells"],
                "has_output_workbook": False, "has_final_answer_tag": False,
                "output_format_valid": False, "fail_reason": "missing_output_workbook"}
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    try:
        workbook = load_workbook(output, data_only=True)
        workbook.close()
    except (zipfile.BadZipFile, InvalidFileException, ParseError, KeyError, ValueError):
        return {"score": 0.0, "matched_cells": 0, "total_cells": case["total_cells"],
                "has_output_workbook": True, "has_final_answer_tag": False,
                "output_format_valid": False, "fail_reason": "invalid_output_workbook"}
    score, matched, total = spreadsheet.score_workbook(output, _local(root, case["reference_path"]),
                                                       case["answer_sheet"], case["answer_position"])
    if total <= 0 or total != case["total_cells"]:
        raise IntegrityError("Scorer region is empty or changed")
    return {"score": score, "matched_cells": matched, "total_cells": total, "has_output_workbook": True,
            "has_final_answer_tag": True, "output_format_valid": True,
            "fail_reason": "" if score else "wrong_answer"}


def _invoke(root, protocol, payload, system, user, mode):
    from wikiskill.isolated.runtime import execute
    _verify_protocol(root)
    return Path(execute(payload, system, user, mode, protocol["timeout_seconds"],
                        libreoffice_app=Path(protocol["libreoffice_app"]),
                        model=protocol["model"], effort=protocol["effort"]))


def _episode(root, protocol, case, arm, skill, directory):
    if case["split"] not in ("train", "val") or (case["split"] == "train" and arm != "s0"):
        raise IntegrityError("Only S0 TRAIN and S0/candidate VAL are authorized")
    expected = {"schema_version": SCHEMA, "uid": case["uid"], "case_id": case["uid"], "split": case["split"],
                "arm": arm, "model": protocol["model"], "reported_model": protocol["model"],
                "reasoning_effort": protocol["effort"], "protocol_sha256": _sha(root / "protocol.json"),
                "skill_sha256": _text_sha(skill), "workspace": str(directory), "attempt_number": 1, "retry_of": None}
    result_path = directory / "result.json"
    if result_path.exists():
        row = _check_record(result_path)
        _audit(directory / "runtime", "spreadsheet", protocol)
        if any(row.get(key) != value for key, value in expected.items()) or any(
                row.get(key) != value for key, value in _score(root, case, directory / "runtime").items()):
            raise IntegrityError("Sealed episode identity or score differs")
        if row["predicted"] != (directory / "runtime/final.txt").read_bytes().decode("utf-8"):
            raise IntegrityError("Sealed prediction differs from native text")
        _check_tree(directory / "input", _read(directory / "input-manifest.json"))
        return row
    payload = directory / "input"
    if payload.exists():
        _check_tree(payload, _read(directory / "input-manifest.json"))
    else:
        payload.mkdir(parents=True)
        shutil.copy2(_local(root, case["input_path"]), payload / "input.xlsx")
        _save(directory / "input-manifest.json", _tree(payload))
    system = (root / "frozen/prompts/spreadsheet.paper.md").read_text().replace("{skill_section}", skill)
    user = json.dumps({"working_directory": "__WORKSPACE__", "instruction": case["question"],
                       "spreadsheet_path": "__WORKSPACE__/input.xlsx", "spreadsheet_content": case["preview"],
                       "instruction_type": case["instruction_type"], "answer_position": case["answer_position"],
                       "answer_sheet": case["answer_sheet"], "output_path": "__WORKSPACE__/output.xlsx"}, ensure_ascii=False)
    user += "\nRuntime: python3 includes openpyxl. soffice is available for formula recalculation. Use a workspace-local profile and output directory; network access is disabled.\n"
    _save(directory / "attempt.json", {**expected, "system_sha256": _text_sha(system), "user_sha256": _text_sha(user)})
    archive = _invoke(root, protocol, payload, system, user, "spreadsheet")
    if archive.resolve() != (directory / "runtime").resolve():
        raise IntegrityError("Unexpected native archive location")
    native = _native(archive)
    row = {**expected, **_score(root, case, archive), "predicted": (archive / "final.txt").read_bytes().decode("utf-8"),
           "returncode": 0, "started_at": native["started_at"], "finished_at": native["finished_at"],
           "system_sha256": _text_sha(system), "user_sha256": _text_sha(user), "artifacts": _tree(archive)}
    _record(result_path, row)
    return row


def _batch(root, protocol, cases, arm, skill, directory):
    """Bound dispatch; on error drain in-flight calls without replenishing."""
    pending_cases = iter(cases)
    completed, errors = {}, []
    with ThreadPoolExecutor(max_workers=protocol["workers"]) as executor:
        futures = {}
        def dispatch():
            case = next(pending_cases, None)
            if case is not None:
                future = executor.submit(_episode, root, protocol, case, arm, skill, directory / case["uid"])
                futures[future] = case["uid"]
        for _ in range(protocol["workers"]):
            dispatch()
        while futures:
            done, _ = wait(futures, return_when=FIRST_COMPLETED)
            for future in done:
                uid = futures.pop(future)
                try:
                    completed[uid] = future.result()
                except BaseException as exc:
                    errors.append(exc)
            if not errors:
                for _ in done:
                    dispatch()
    if errors:
        raise StudyError(f"Batch stopped with preserved attempts: {type(errors[0]).__name__}: {errors[0]}") from errors[0]
    rows = [completed[case["uid"]] for case in cases]
    _save(directory / "outcomes.json", rows)
    return rows


ROLE_STATE = """
The Wiki is research notes; inference receives only the task, input workbook,
permitted tools and accepted SKILL.md. It cannot read this Wiki or references.
skills.json is the active skill authority and is empty in this first round.
An already documented Wiki pattern may justify the first skill when it adds
useful actions beyond the inference prompt. no_action remains valid: do not
force a candidate. If returning no_action, explain insufficient evidence,
already-deployed guidance, or already-fixed task interfaces before finish.
Read role-manifest.json and use its exact paths. trace-summary.json lists every
available TRAIN execution trace and its training-reference JSON. Read paths
verbatim; do not guess or probe filenames. Large references have an exact
target_cell_pages index; all target cells are retained in those page files.
Reference values belong to TRAIN only. Diagnose failures using them but never
encode task IDs, answer values, private paths or task-specific facts in skills.
No validation/test tasks, traces, scores or references are supplied to roles.
The observed scores are unchanged. A valid output workbook is scored without
requiring any particular final-text tag or explanation. Tool setup, task fields
and scorer are identical for fresh TRAIN and both validation arms.
"""


def _feedback(root, row, case):
    if row["split"] != "train" or case["split"] != "train" or row["uid"] != case["uid"]:
        raise IntegrityError("Only matching TRAIN data can enter role feedback")
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    gold = load_workbook(_local(root, case["reference_path"]), data_only=True)
    output_path = Path(row["workspace"]) / "runtime/output.xlsx"
    output = load_workbook(output_path, data_only=True) if row["output_format_valid"] else None
    formulas = load_workbook(output_path, data_only=False) if output else None
    cells = []
    try:
        for name, position in spreadsheet.parse_answer_regions(case["answer_sheet"], case["answer_position"]):
            expected = gold[name] if name else gold.worksheets[0]
            actual = (output[name] if name in output.sheetnames else None) if output and name else output.worksheets[0] if output else None
            formula_sheet = formulas[actual.title] if actual is not None else None
            for col, index in spreadsheet._region(position, expected.max_row):
                address = f"{get_column_letter(col)}{index}"
                reference = expected[address].value
                predicted = actual[address].value if actual is not None else None
                cells.append({"sheet": expected.title, "cell": address, "predicted": _plain(predicted),
                              "reference": _plain(reference), "match": actual is not None and spreadsheet._norm(predicted) == spreadsheet._norm(reference),
                              "predicted_formula_or_value": _plain(formula_sheet[address].value) if formula_sheet is not None else None})
    finally:
        gold.close()
        if output:
            output.close()
            formulas.close()
    if (sum(cell["match"] for cell in cells), len(cells)) != (row["matched_cells"], row["total_cells"]):
        raise IntegrityError("TRAIN feedback does not reproduce sealed scorer counts")
    return {"uid": row["uid"], "split": "train", "original_score": row["score"],
            "predicted_final_response": row["predicted"], "answer_sheet": case["answer_sheet"],
            "answer_position": case["answer_position"], "matched_cells": row["matched_cells"],
            "total_cells": row["total_cells"], "target_cells": cells}


def _stage_feedback(payload, relative, feedback):
    cells = feedback["target_cells"]
    if len(_json(feedback)) <= 64000:
        _save(payload / relative, feedback)
        return
    pages = []
    for start in range(0, len(cells), 200):
        part = start // 200 + 1
        path = f"training-reference-parts/{feedback['uid']}/part-{part:04}.json"
        chunk = cells[start:start + 200]
        _save(payload / path, {"uid": feedback["uid"], "split": "train", "target_cells": chunk})
        pages.append({"path": path, "cell_count": len(chunk), "mismatch_count": sum(not cell["match"] for cell in chunk)})
    _save(payload / relative, {**{key: value for key, value in feedback.items() if key != "target_cells"},
                               "target_preview": cells[:4], "target_cell_pages": pages,
                               "complete_values": "All values are retained in target_cell_pages; preview is incomplete."})


def _role_payload(root, protocol, directory, role, rows, cases, wiki):
    if role not in ("maintainer", "proposer") or any(row["split"] != "train" for row in rows):
        raise IntegrityError("Only learning roles can receive TRAIN references")
    payload = directory / "input"
    if payload.exists():
        _check_tree(payload, _read(directory / "input-manifest.json"))
        return payload
    payload.mkdir(parents=True)
    chosen = evidence.sample(rows, protocol["sampling"]["seed"])
    selected = chosen if role == "maintainer" else rows
    by = {case["uid"]: case for case in cases}
    summary = []
    for row in selected:
        case = by[row["uid"]]
        feedback = _feedback(root, row, case)
        trace_path = f"traces/{row['uid']}.md"
        reference_path = f"training-reference/{row['uid']}.json"
        _save_text(payload / trace_path, evidence.visible_trace(row, case, limit=15000))
        _stage_feedback(payload, reference_path, feedback)
        summary.append({"uid": row["uid"], "split": "train", "outcome": "PASS" if row["score"] else "FAIL",
                        "original_score": row["score"], "path": trace_path, "feedback_path": reference_path,
                        "prediction": row["predicted"], "matched_cells": row["matched_cells"], "total_cells": row["total_cells"]})
    for relative, text in wiki.items():
        _save_text(_local(payload / "wiki", relative), text)
    _save(payload / "skills.json", {})
    _save(payload / "trace-summary.json", summary)
    _save(payload / "role-manifest.json", {"role": role, "data_split": "train", "wiki_index": "wiki/index.md",
                                            "trial_history": "wiki/skill-impact.md", "active_skills": "skills.json",
                                            "active_skill_names": [], "wiki_is_deployed": False,
                                            "trace_summary": "trace-summary.json", "trace_paths": [item["path"] for item in summary],
                                            "reference_paths": [item["feedback_path"] for item in summary]})
    _save(directory / "sample.json", {"seed": protocol["sampling"]["seed"], "selected": [row["uid"] for row in chosen],
                                       "failures": sum(row["score"] == 0 for row in chosen), "successes": sum(row["score"] == 1 for row in chosen),
                                       "available_failures": sum(row["score"] == 0 for row in rows), "available_successes": sum(row["score"] == 1 for row in rows)})
    _save(directory / "input-manifest.json", _tree(payload))
    return payload


def _trace_reads(archive, payload):
    allowed = {item["path"]: item["uid"] for item in _read(payload / "trace-summary.json")}
    seen = set()
    path = archive / "tool-events.jsonl"
    if path.exists():
        for event in read_jsonl(path):
            if event.get("tool") == "read_file" and event.get("ok") is True:
                relative = event.get("arguments", {}).get("path", "")
                if relative in allowed:
                    seen.add(allowed[relative])
    return sorted(seen)


def _learn(root, protocol, rows, cases):
    wiki = _read(root / "frozen/initial-wiki.json")
    for role in ("maintainer", "proposer"):
        directory = root / "learning" / role
        payload = _role_payload(root, protocol, directory, role, rows, cases, wiki)
        system = (root / "frozen/prompts" / f"{role}.paper.md").read_text().replace("{task_desc}", "spreadsheet manipulation tasks with isolated bash")
        system += "\nSubmission transport: read_file reads staged files; finish(proposal) submits the exact paper JSON. No extra computation or network tool is available.\n" + ROLE_STATE
        user = "Read the exact manifest, Wiki and TRAIN trace summary. Use read_file and finish.\n"
        if role == "maintainer":
            user += "Analyze the sampled real TRAIN trajectories below (failure/success classes were sampled separately).\n"
            for item in _read(payload / "trace-summary.json"):
                user += (payload / item["path"]).read_text() + "\n"
        else:
            user += "Read at least four distinct actual execution trace files before creating or patching a skill.\n"
        user += "\nEXACT ROLE FILE MANIFEST:\n" + _json(_read(payload / "role-manifest.json"))
        user += "\nTRAIN OUTCOME / PREDICTION / REFERENCE SUMMARY:\n" + _json(_read(payload / "trace-summary.json"))
        _save(directory / "attempt.json", {"mode": role, "system_sha256": _text_sha(system), "user_sha256": _text_sha(user),
                                           "protocol_sha256": _sha(root / "protocol.json")})
        _state(root, role)
        archive = _invoke(root, protocol, payload, system, user, role)
        _native(archive)
        submitted = _read(archive / "submission.json")
        if role == "maintainer":
            normalized, wiki = contracts.wiki_update(submitted["proposal"], wiki)
            wiki["log.md"] += "\n## Iteration 1\n\n" + normalized["append_log"] + "\n"
            result = {"change": normalized, "wiki": wiki, "artifacts": _tree(archive)}
            _record(directory / "result.json", result)
            _save(root / "wiki-after-maintainer.json", wiki)
        else:
            reads = _trace_reads(archive, payload)
            if reads != sorted(set(submitted.get("read_trace_ids", []))):
                raise IntegrityError("Proposer trace IDs do not match successful actual file reads")
            change, candidate = contracts.proposal(submitted["proposal"], {}, reads)
            result = {"change": change, "candidate_skills": candidate, "read_trace_ids": reads, "artifacts": _tree(archive)}
            _record(directory / "result.json", result)
            _save(root / "proposal.json", {key: value for key, value in result.items() if key != "artifacts"})
    return result["change"], result["candidate_skills"]


def _all_artifacts(root):
    excluded = {"study.lock", "state.json", "state.json.tmp", "final-freeze.json"}
    return {relative: digest for relative, digest in _tree(root).items()
            if relative not in excluded and not relative.startswith("events/")}


def run(root: Path) -> dict:
    """Run/resume exactly one prepared study. Native errors never auto-retry."""
    root = Path(root).resolve()
    with _lock(root):
        try:
            protocol = _verify_protocol(root)
            if (root / "final-freeze.json").exists():
                verify(root)
                if status(root)["phase"] != "complete":
                    _state(root, "complete", gate=_read(root / "gate.json"), next_stage="none")
                return status(root)
            cases = _read(root / "frozen/cases.json")
            _state(root, "fresh_s0_val")
            baseline = _batch(root, protocol, cases["val"], "s0", "", root / "baseline")
            _state(root, "fresh_s0_train")
            train = _batch(root, protocol, cases["train"], "s0", "", root / "train")
            change, candidate = _learn(root, protocol, train, cases["train"])
            candidate_rows = None
            if change["action"] != "no_action":
                _state(root, "candidate_val")
                candidate_rows = _batch(root, protocol, cases["val"], "candidate", contracts.skill_text(candidate), root / "candidate")
            baseline_correct = sum(row["score"] for row in baseline)
            candidate_correct = sum(row["score"] for row in candidate_rows) if candidate_rows is not None else None
            accepted = candidate_correct is not None and candidate_correct > baseline_correct
            active = candidate if accepted else {}
            gate = {"schema_version": SCHEMA, "iteration": 1, "action": change["action"], "accepted": accepted,
                    "verdict": "ACCEPT" if accepted else "NO_ACTION" if candidate_rows is None else "REJECT",
                    "baseline_correct": baseline_correct, "candidate_correct": candidate_correct, "val_n": len(baseline),
                    "active_skills": active, "protocol_sha256": _sha(root / "protocol.json"), "test_started": False,
                    "scope": protocol["scope"], "next_stage": "none"}
            _save(root / "gate.json", gate)
            _save(root / "FINAL-SKILLS.json", active)
            _save_text(root / "FINAL-SKILL.md", contracts.skill_text(active))
            _save(root / "skill-impact.json", {"iteration": 1, "proposal": change, "gate": gate,
                                                "candidate_skills": candidate, "wiki_retained": True})
            _save(root / "final-freeze.json", {"schema_version": SCHEMA, "protocol_sha256": _sha(root / "protocol.json"),
                                               "artifacts": _all_artifacts(root), "gate": gate})
            _state(root, "complete", gate=gate, next_stage="none")
            verify(root)
            return status(root)
        except BaseException as exc:
            _state(root, "needs_attention", error=f"{type(exc).__name__}: {exc}", next_stage="none",
                   recovery="Resume only unchanged inputs and sealed native calls; incomplete attempts require a separate study")
            raise


def verify(root: Path) -> dict:
    """Read-only verification of frozen inputs, native seals and all scores."""
    root = Path(root).resolve()
    protocol = _verify_protocol(root)
    cases = _read(root / "frozen/cases.json")
    by = {split: {case["uid"]: case for case in cases[split]} for split in ("train", "val")}
    record_count = 0
    for stage, split in (("baseline", "val"), ("train", "train"), ("candidate", "val")):
        for archive in sorted((root / stage).glob("*/runtime")):
            if not (archive.parent / "result.json").exists():
                _native(archive)
                _audit(archive, "spreadsheet", protocol)
        for path in sorted((root / stage).glob("*/result.json")):
            row = _check_record(path)
            _audit(path.parent / "runtime", "spreadsheet", protocol)
            if row["split"] != split or row["uid"] not in by[split]:
                raise IntegrityError("Record outside frozen split")
            score = _score(root, by[split][row["uid"]], path.parent / "runtime")
            if any(row.get(key) != value for key, value in score.items()):
                raise IntegrityError("Recomputed score differs from frozen record")
            _check_tree(path.parent / "input", _read(path.parent / "input-manifest.json"))
            record_count += 1
    for archive in sorted((root / "learning").glob("*/runtime")):
        if not (archive.parent / "result.json").exists():
            _native(archive)
            _audit(archive, archive.parent.name, protocol)
    for path in sorted((root / "learning").glob("*/result.json")):
        _check_record(path)
        _audit(path.parent / "runtime", path.parent.name, protocol)
        _check_tree(path.parent / "input", _read(path.parent / "input-manifest.json"))
    if (root / "final-freeze.json").exists():
        frozen = _read(root / "final-freeze.json")
        if frozen["artifacts"] != _all_artifacts(root):
            raise IntegrityError("Final frozen artifact membership or hashes changed")
        gate = _read(root / "gate.json")
        baseline = _read(root / "baseline/outcomes.json")
        candidates = _read(root / "candidate/outcomes.json") if gate["candidate_correct"] is not None else None
        if len(baseline) != len(by["val"]) or {row["uid"] for row in baseline} != set(by["val"]):
            raise IntegrityError("Incomplete frozen baseline")
        if candidates is not None and (len(candidates) != len(baseline) or {row["uid"] for row in candidates} != set(by["val"])):
            raise IntegrityError("Incomplete frozen candidate")
        expected_baseline = sum(row["score"] for row in baseline)
        expected_candidate = sum(row["score"] for row in candidates) if candidates is not None else None
        if gate["baseline_correct"] != expected_baseline or gate["candidate_correct"] != expected_candidate or gate["accepted"] != (expected_candidate is not None and expected_candidate > expected_baseline):
            raise IntegrityError("Frozen gate does not recompute")
    return {"verified": True, "schema_version": SCHEMA, "protocol_sha256": _sha(root / "protocol.json"),
            "score_records": record_count, "complete": (root / "final-freeze.json").exists(), "budget": protocol["budget"]}


def status(root: Path) -> dict:
    """Read-only operational snapshot; call verify for integrity validation."""
    root = Path(root).resolve()
    value = _read(root / "state.json") if (root / "state.json").exists() else {"phase": "not_prepared", "active": False}
    return {**value, "root": str(root), "complete": (root / "final-freeze.json").exists(),
            "score_records": {stage: len(list((root / stage).glob("*/result.json"))) for stage in ("baseline", "train", "candidate")},
            "gate": _read(root / "gate.json") if (root / "gate.json").exists() else None}
