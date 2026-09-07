"""Offline end-to-end checks of the bounded study, with no model calls."""
from __future__ import annotations

import json
from pathlib import Path
import shutil
import sys
import types

from openpyxl import Workbook, load_workbook
import pytest

from wikiskill.spreadsheet import study


@pytest.fixture(autouse=True)
def codex_binding(monkeypatch):
    def binding(*, with_version=True):
        result = {"path": "/synthetic/codex", "sha256": "synthetic"}
        if with_version:
            result["version"] = "codex synthetic"
        return result
    monkeypatch.setattr(study, "_codex_binding", binding)


@pytest.fixture
def data(tmp_path):
    data = tmp_path / "data"
    data.mkdir()
    splits = tmp_path / "splits"
    splits.mkdir()
    items = []
    for number in range(1, 13):
        uid = f"{number:03}"
        directory = data / uid
        directory.mkdir()
        for kind, value in (("init", 0), ("golden", 90000 + number)):
            workbook = Workbook()
            workbook.active["A1"] = value
            workbook.save(directory / f"task_{kind}.xlsx")
            workbook.close()
        items.append({"id": uid, "spreadsheet_path": uid, "instruction": "Fill the requested result cell.",
                      "instruction_type": "Cell-Level Manipulation", "answer_position": "A1", "answer_sheet": "Sheet"})
    (data / "dataset.json").write_text(json.dumps(items))
    # Deliberately reverse input order: study selection must be UID-stable.
    (splits / "train.json").write_text(json.dumps([{"uid": f"{number:03}"} for number in range(8, 0, -1)]))
    (splits / "val.json").write_text(json.dumps([{"uid": f"{number:03}"} for number in range(12, 8, -1)]))
    (splits / "test.json").write_text("TEST MUST NEVER BE PARSED")
    app = tmp_path / "LibreOffice.app"
    binary = app / "Contents/MacOS/soffice"
    binary.parent.mkdir(parents=True)
    binary.write_text("synthetic binary for offline tests")
    return data, splits, app


@pytest.fixture
def fake_runtime(monkeypatch):
    """Model simulator uses reference data only in test code, outside payload."""
    runtime = types.ModuleType("wikiskill.isolated.runtime")
    state = types.SimpleNamespace(calls=[], audits=[], action="create", baseline_correct=1,
                                  candidate_correct=4, all_train_pass=False, fail=False,
                                  missing_train_output=False, summary_only=False)
    def execute(payload, system, user, mode, timeout=1800, *, libreoffice_app, model, effort):
        payload = Path(payload)
        archive = payload.parent / "runtime"
        if (archive / "native-complete.json").exists():
            return archive
        if archive.exists():
            raise study.IntegrityError("Incomplete native call preserved; no retry")
        archive.mkdir()
        state.calls.append((mode, payload, system, user))
        (archive / "request.json").write_text(json.dumps({"mode": mode, "system": system, "user_template": user}))
        if state.fail:
            (archive / "failure.json").write_text('{"error":"synthetic timeout"}')
            raise TimeoutError("synthetic timeout")
        (archive / "final.txt").write_text("")  # No final text/tag is needed for workbook scoring.
        receipts = []
        if mode == "spreadsheet":
            receipts = [{"at": 1, "tool": "bash", "arguments": {"command": "python3 edit.py"}, "ok": True, "result": "Saved output workbook."}]
            assert {p.name for p in payload.iterdir()} == {"input.xlsx"}
            assert "reference" not in user and "golden" not in user
            task = json.loads(user.split("\nRuntime:", 1)[0])
            assert task["instruction_type"] == "Cell-Level Manipulation"
            assert task["answer_position"] == "A1" and task["answer_sheet"] == "Sheet"
            assert task["spreadsheet_content"]["Sheet"][0][0]["value"] == 0
            stage = payload.parent.parent.name
            number = int(payload.parent.name)
            correct = ((number - 9 < state.baseline_correct) if stage == "baseline" else
                       (number - 9 < state.candidate_correct) if stage == "candidate" else
                       state.all_train_pass or number % 2 == 0)
            if not (stage == "train" and state.missing_train_output):
                workbook = load_workbook(payload / "input.xlsx")
                workbook.active["A1"] = 90000 + number if correct else -1
                workbook.save(archive / "output.xlsx")
                workbook.close()
        else:
            summary = json.loads((payload / "trace-summary.json").read_text())
            assert all(item["split"] == "train" and int(item["uid"]) <= 8 for item in summary)
            assert not list(payload.rglob("*.xlsx"))
            assert "EXACT ROLE FILE MANIFEST" in user
            assert "trace-summary.json" in system and "training-reference" in system
            if mode == "maintainer":
                proposal = {"create_patterns": [{"name": "Verify cached output", "content": "# Verify cached output\nRead back saved workbook cells before finishing.\n"}],
                            "update_patterns": [],
                            "update_index": "# Patterns\n- [Verify cached output](wiki/patterns/Verify cached output.md): Read back saved cells.\n",
                            "append_log": "Observed both successful and failed saved outputs."}
                reads = []
            else:
                reads = [item["uid"] for item in summary[:4]]
                if not state.summary_only:
                    receipts = [{"at": 1, "tool": "read_file", "arguments": {"path": item["path"]},
                                 "ok": True, "result": "actual trace\u2028still one JSON record\u0085"} for item in summary[:4]]
                proposal = {"action": "no_action"} if state.action == "no_action" else {
                    "action": "create", "name": "verify_output",
                    "skill_md": "---\nname: verify_output\ndescription: Verify spreadsheet edits\n---\n## When to Apply\nAfter edits.\n## When NOT to Apply\nRead-only tasks.\n## Instructions\nRead back saved workbook target cells.\n",
                    "purpose_md": "# Origin\nTraining traces.\n# Patterns Addressed\nVerify cached output.\n# Evolution History\nInitial candidate.\n"}
            (archive / "submission.json").write_text(json.dumps({"proposal": proposal, "read_trace_ids": reads}))
        (archive / "tool-events.jsonl").write_text("".join(json.dumps(event, ensure_ascii=False) + "\n" for event in receipts))
        artifacts = study._tree(archive)
        (archive / "native-complete.json").write_text(json.dumps({"started_at": "2026-09-07T00:00:00Z", "finished_at": "2026-09-07T00:01:00Z", "duration": 60, "artifacts": artifacts}))
        return archive
    def audit(archive, mode, **kwargs):
        state.audits.append((str(archive), mode, kwargs))
        return {"mode": mode}
    runtime.execute = execute
    runtime.audit = audit
    runtime.verify_native_complete = study._native
    monkeypatch.setitem(sys.modules, "wikiskill.isolated.runtime", runtime)
    return state


def prepared(tmp_path, data, **kwargs):
    root = tmp_path / "study"
    data_path, splits, app = data
    study.prepare(root, data=data_path, split_dir=splits, libreoffice_app=app, workers=1, **kwargs)
    return root


def test_prepare_freezes_uid_prefixes_and_inputs_without_test(tmp_path, data):
    root = prepared(tmp_path, data, train_limit=4, val_limit=2)
    protocol = study._read(root / "protocol.json")
    assert protocol["selection"]["train"] == ["001", "002", "003", "004"]
    assert protocol["selection"]["val"] == ["009", "010"]
    assert protocol["budget"]["inference_max"] == 8
    assert not any("test" in path for path in protocol["frozen_files"])
    assert study.verify(root)["score_records"] == 0
    # Frozen snapshots remain usable even if the external dataset is moved.
    shutil.rmtree(data[0])
    assert study.verify(root)["verified"]
    with pytest.raises(study.IntegrityError, match="empty"):
        prepared(tmp_path, data)


def test_complete_accepts_only_strict_gain_and_freezes_once(tmp_path, data, fake_runtime):
    root = prepared(tmp_path, data)
    result = study.run(root)
    assert result["phase"] == "complete"
    assert result["gate"]["verdict"] == "ACCEPT"
    assert result["gate"]["baseline_correct"] == 1
    assert result["gate"]["candidate_correct"] == 4
    assert len(fake_runtime.calls) == 18
    assert [call[0] for call in fake_runtime.calls] == ["spreadsheet"] * 12 + ["maintainer", "proposer"] + ["spreadsheet"] * 4
    train_sample = study._read(root / "learning/maintainer/sample.json")
    assert (train_sample["failures"], train_sample["successes"]) == (4, 3)
    baseline_row = study._read(root / "baseline/009/result.json")
    assert baseline_row["predicted"] == "" and baseline_row["score"] == 1
    assert baseline_row["has_final_answer_tag"] is True
    assert study.verify(root)["score_records"] == 16
    before = study._sha(root / "final-freeze.json")
    assert study.run(root)["complete"] is True
    assert len(fake_runtime.calls) == 18 and study._sha(root / "final-freeze.json") == before
    assert fake_runtime.audits


def test_tie_rejects_and_keeps_rejected_candidate(tmp_path, data, fake_runtime):
    fake_runtime.candidate_correct = 1
    root = prepared(tmp_path, data)
    result = study.run(root)
    assert result["gate"]["verdict"] == "REJECT"
    assert (root / "FINAL-SKILL.md").read_text() == ""
    assert study._read(root / "skill-impact.json")["candidate_skills"]
    assert study._read(root / "wiki-after-maintainer.json")["patterns/Verify cached output.md"]


def test_no_action_class_shortage_stops_without_candidate_or_resampling(tmp_path, data, fake_runtime):
    fake_runtime.action = "no_action"
    fake_runtime.all_train_pass = True
    root = prepared(tmp_path, data)
    result = study.run(root)
    assert result["gate"]["verdict"] == "NO_ACTION"
    assert len(fake_runtime.calls) == 14
    sample = study._read(root / "learning/maintainer/sample.json")
    assert sample["failures"] == 0 and sample["successes"] == 3
    assert sample["available_successes"] == 8
    assert len(study._read(root / "learning/proposer/input/trace-summary.json")) == 8
    assert not (root / "candidate").exists()


def test_missing_output_is_scored_zero_and_reference_feedback_matches(tmp_path, data, fake_runtime):
    fake_runtime.action = "no_action"
    fake_runtime.missing_train_output = True
    root = prepared(tmp_path, data)
    study.run(root)
    row = study._read(root / "train/001/result.json")
    assert row["score"] == 0 and row["total_cells"] == 1
    assert row["has_output_workbook"] is False and row["has_final_answer_tag"] is False
    assert row["fail_reason"] == "missing_output_workbook"
    assert study._read(root / "learning/maintainer/sample.json")["failures"] == 5


def test_failed_native_call_stops_without_score_or_retry(tmp_path, data, fake_runtime):
    root = prepared(tmp_path, data)
    fake_runtime.fail = True
    with pytest.raises(study.StudyError, match="synthetic timeout"):
        study.run(root)
    assert study.status(root)["phase"] == "needs_attention"
    assert len(fake_runtime.calls) == 1
    assert not list(root.rglob("result.json"))
    assert (root / "baseline/009/runtime/failure.json").exists()
    with pytest.raises(study.IntegrityError, match="Incomplete native"):
        study.verify(root)
    fake_runtime.fail = False
    with pytest.raises(study.StudyError, match="no retry"):
        study.run(root)
    assert len(fake_runtime.calls) == 1
    assert not (root / "gate.json").exists()


def test_sealed_native_call_can_resume_deterministic_postprocessing(tmp_path, data, fake_runtime, monkeypatch):
    root = prepared(tmp_path, data)
    fake_runtime.action = "no_action"
    original = study._score
    failed = False
    def interrupted(*args):
        nonlocal failed
        if not failed:
            failed = True
            raise OSError("synthetic scoring interruption")
        return original(*args)
    monkeypatch.setattr(study, "_score", interrupted)
    with pytest.raises(study.StudyError, match="scoring interruption"):
        study.run(root)
    assert (root / "baseline/009/runtime/native-complete.json").exists()
    assert not (root / "baseline/009/result.json").exists()
    assert study.verify(root)["verified"]  # Sealed raw data can await deterministic scoring.
    assert study.run(root)["complete"]
    assert len(fake_runtime.calls) == 14


def test_proposer_must_read_real_traces_not_only_summary(tmp_path, data, fake_runtime):
    root = prepared(tmp_path, data)
    fake_runtime.summary_only = True
    with pytest.raises(study.IntegrityError, match="actual file reads"):
        study.run(root)
    assert study.status(root)["phase"] == "needs_attention"
    assert not (root / "gate.json").exists()
    assert (root / "learning/proposer/runtime/submission.json").exists()


def test_verify_rejects_frozen_input_and_final_artifact_tampering(tmp_path, data, fake_runtime):
    root = prepared(tmp_path, data)
    fake_runtime.action = "no_action"
    study.run(root)
    target = root / "learning/proposer/input/traces/001.md"
    target.write_text("changed trace")
    with pytest.raises(study.IntegrityError, match="tree"):
        study.verify(root)


def test_prepare_rejects_overlap_limits_and_model_before_dispatch(tmp_path, data):
    with pytest.raises(ValueError, match="train_limit"):
        prepared(tmp_path, data, train_limit=9)
    with pytest.raises(ValueError, match="luna/high"):
        prepared(tmp_path, data, model="different-model")
    (data[1] / "val.json").write_text('[{"uid":"001"}]')
    with pytest.raises(study.IntegrityError, match="overlap"):
        prepared(tmp_path, data)


def test_exclusive_writer_lock_prevents_second_controller(tmp_path, data, fake_runtime):
    root = prepared(tmp_path, data)
    with study._lock(root):
        with pytest.raises(study.StudyError, match="writer lock"):
            study.run(root)
    assert not fake_runtime.calls


def test_completed_freeze_recovers_interrupted_status_projection(tmp_path, data, fake_runtime, monkeypatch):
    root = prepared(tmp_path, data)
    fake_runtime.action = "no_action"
    original = study._state
    interrupted = False
    def state(directory, phase, **fields):
        nonlocal interrupted
        if phase == "complete" and not interrupted:
            interrupted = True
            raise OSError("synthetic status interruption")
        return original(directory, phase, **fields)
    monkeypatch.setattr(study, "_state", state)
    with pytest.raises(OSError, match="status interruption"):
        study.run(root)
    assert study.status(root)["phase"] == "needs_attention"
    assert study.run(root)["phase"] == "complete"
    assert len(fake_runtime.calls) == 14
