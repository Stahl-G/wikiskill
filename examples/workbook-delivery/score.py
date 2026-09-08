"""Score the exact delivered workbook; no model or recalculation process is run.

The task uses a narrow, documented formula vocabulary: multiplication/PRODUCT
for each amount and SUM/addition for the total. Expected values come from the
unaltered input. This checker is a teaching example, not a general Excel judge.
"""
import json
import math
import re
import sys
from pathlib import Path
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import load_workbook


def score(payload):
    task = payload['task']
    source = Path(task['files'][0])
    # Source/configuration failures propagate as scorer failures, never model zeros.
    original = load_workbook(source, data_only=False)
    end = task['input']['total_row']
    rows = range(4, end)
    targets = {f'D{r}' for r in range(4, end + 1)}
    sheet = original['Order']
    amounts = {f'D{r}': sheet[f'B{r}'].value * sheet[f'C{r}'].value for r in rows}
    amounts[f'D{end}'] = sum(amounts.values())
    try:
        formula_book = load_workbook(payload['output']['path'], data_only=False)
        value_book = load_workbook(payload['output']['path'], data_only=True)
    except (BadZipFile, ValueError, KeyError, ParseError, InvalidFileException):
        return {'score': 0, 'success': False, 'feedback': 'The delivered file is not a readable XLSX workbook.'}
    problems = []
    if formula_book.sheetnames != original.sheetnames:
        return {'score': 0, 'success': False, 'feedback': 'Preserve the original worksheet names and order.'}
    output, cached = formula_book['Order'], value_book['Order']
    formulas_ok = True
    for coord in sorted(targets):
        raw = output[coord].value
        expression = re.sub(r'\s|\$', '', raw.upper()) if isinstance(raw, str) else ''
        row = int(coord[1:])
        if row < end:
            alternatives = {f'=B{row}*C{row}', f'=C{row}*B{row}', f'=PRODUCT(B{row},C{row})', f'=PRODUCT(B{row}:C{row})'}
        else:
            alternatives = {f'=SUM(D4:D{end-1})', '=' + '+'.join(f'D{r}' for r in rows), '=SUM(' + ','.join(f'D{r}' for r in rows) + ')'}
        if expression not in alternatives:
            formulas_ok = False
            problems.append(f'{coord}: use the requested live multiplication/PRODUCT or SUM formula.')
    values_ok = True
    for coord, expected in amounts.items():
        actual = cached[coord].value
        if isinstance(actual, bool) or not isinstance(actual, (int, float)) or not math.isclose(actual, expected, rel_tol=0, abs_tol=1e-8):
            values_ok = False
            problems.append(f'{coord}: saved calculated value is missing or incorrect; check the delivered file after recalculation.')
    preserved = True
    # Compare values outside the authorized target range, plus visible number formats everywhere.
    for r in range(1, max(output.max_row, sheet.max_row) + 1):
        for c in range(1, max(output.max_column, sheet.max_column) + 1):
            old, new = sheet.cell(r, c), output.cell(r, c)
            if (old.coordinate not in targets and old.value != new.value) or old.number_format != new.number_format:
                preserved = False
    if not preserved:
        problems.append('Unrelated values or number formats changed.')
    passed = sum((formulas_ok, values_ok, preserved))
    for book in (original, formula_book, value_book):
        book.close()
    return {'score': passed / 3, 'success': passed == 3,
            'feedback': 'All three checks passed: formulas, saved values, preserved inputs/number formats.' if passed == 3 else ' '.join(problems)}


if __name__ == '__main__':
    print(json.dumps(score(json.load(sys.stdin))))
